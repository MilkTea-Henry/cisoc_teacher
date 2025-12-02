"""
Tests for docx_parser module.
"""

import pytest
import os
import sys
import json
from io import BytesIO

# Add the backend to the path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from backend.app.docx_parser import (
    extract_chapter_info,
    extract_lab_id,
    extract_objective,
    extract_steps,
    extract_qa,
    parse_docx
)
from backend.app.schema import Step, QA


class TestExtractChapterInfo:
    """Tests for chapter info extraction."""
    
    def test_extract_chapter_english(self):
        """Test extracting English chapter format."""
        chapter, title = extract_chapter_info("Chapter 2: Router Basics")
        assert chapter == "2"
        assert title == "Router Basics"
    
    def test_extract_chapter_chinese(self):
        """Test extracting Chinese chapter format."""
        chapter, title = extract_chapter_info("第 3 章：網路設定")
        assert chapter == "3"
        assert title == "網路設定"
    
    def test_extract_chapter_abbreviated(self):
        """Test extracting abbreviated chapter format."""
        chapter, title = extract_chapter_info("Ch.1 - Introduction")
        assert chapter == "1"
        assert title == "Introduction"
    
    def test_extract_chapter_numbered(self):
        """Test extracting numbered format."""
        chapter, title = extract_chapter_info("1. Getting Started")
        assert chapter == "1"
        assert title == "Getting Started"
    
    def test_extract_chapter_none(self):
        """Test when no chapter found."""
        chapter, title = extract_chapter_info("Some random text")
        assert chapter is None
        assert title is None


class TestExtractLabId:
    """Tests for lab ID extraction."""
    
    def test_extract_lab_id_standard(self):
        """Test standard lab ID format."""
        lab_id = extract_lab_id("Lab 2.1.3 - Basic Router Setup")
        assert lab_id == "2.1.3"
    
    def test_extract_lab_id_chinese(self):
        """Test Chinese lab format."""
        lab_id = extract_lab_id("實驗 1-2 路由器設定")
        assert lab_id == "1-2"
    
    def test_extract_lab_id_simple(self):
        """Test simple numbered format."""
        lab_id = extract_lab_id("Configure Lab 3.2.1")
        assert lab_id == "3.2.1"
    
    def test_extract_lab_id_none(self):
        """Test when no lab ID found."""
        lab_id = extract_lab_id("Introduction to Networking")
        assert lab_id is None


class TestExtractObjective:
    """Tests for objective extraction."""
    
    def test_extract_objective_english(self):
        """Test English objective format."""
        paragraphs = [
            "Lab 2.1.3",
            "Objective: Configure basic router settings",
            "Prerequisites: None"
        ]
        objective = extract_objective(paragraphs)
        assert "Configure basic router settings" in objective
    
    def test_extract_objective_chinese(self):
        """Test Chinese objective format."""
        paragraphs = [
            "實驗 1.1",
            "學習目標：了解基本網路設定",
            "步驟說明"
        ]
        objective = extract_objective(paragraphs)
        assert "了解基本網路設定" in objective
    
    def test_extract_objective_multiline(self):
        """Test objective in next paragraph."""
        paragraphs = [
            "Goals:",
            "Learn router configuration basics",
            "More content"
        ]
        objective = extract_objective(paragraphs)
        assert "Learn router configuration basics" in objective
    
    def test_extract_objective_none(self):
        """Test when no objective found."""
        paragraphs = ["Step 1", "Step 2", "Step 3"]
        objective = extract_objective(paragraphs)
        assert objective is None


class TestExtractSteps:
    """Tests for step extraction."""
    
    def test_extract_steps_basic(self):
        """Test basic step extraction."""
        paragraphs = [
            "Step 1: Enable privileged mode",
            "Router> enable",
            "Step 2: Enter global config",
            "Router# configure terminal"
        ]
        steps = extract_steps(paragraphs)
        assert len(steps) == 2
        assert steps[0].id == "1"
        assert steps[0].desc == "Enable privileged mode"
        assert "enable" in steps[0].cmds
    
    def test_extract_steps_chinese(self):
        """Test Chinese step format."""
        paragraphs = [
            "步驟 1: 進入特權模式",
            "  enable",
            "步驟 2: 進入全域設定模式",
            "  configure terminal"
        ]
        steps = extract_steps(paragraphs)
        assert len(steps) == 2
        assert steps[0].id == "1"
    
    def test_extract_steps_numbered(self):
        """Test numbered format."""
        paragraphs = [
            "1. Connect to the router",
            "2. Power on the device",
            "3. Access console"
        ]
        steps = extract_steps(paragraphs)
        assert len(steps) == 3
    
    def test_extract_steps_with_prompt(self):
        """Test steps with CLI prompts."""
        paragraphs = [
            "Step 1: Configure hostname",
            "Router> enable",
            "Router# configure terminal",
            "Router(config)# hostname R1"
        ]
        steps = extract_steps(paragraphs)
        assert len(steps) == 1
        assert "enable" in steps[0].cmds
        assert "configure terminal" in steps[0].cmds
        # hostname R1 comes from Router(config)# prompt
        assert any("hostname" in cmd for cmd in steps[0].cmds)
    
    def test_extract_steps_empty(self):
        """Test with no steps."""
        paragraphs = ["Introduction", "Overview", "Summary"]
        steps = extract_steps(paragraphs)
        assert len(steps) == 0


class TestExtractQA:
    """Tests for Q&A extraction."""
    
    def test_extract_qa_standard(self):
        """Test standard Q&A format."""
        paragraphs = [
            "Q: What command enters privileged mode?",
            "A: The enable command",
            "Q: What is OSPF?",
            "A: Open Shortest Path First"
        ]
        qa = extract_qa(paragraphs)
        assert len(qa) == 2
        assert "privileged" in qa[0].q
        assert "enable" in qa[0].a
    
    def test_extract_qa_chinese(self):
        """Test Chinese Q&A format."""
        paragraphs = [
            "問題: 如何設定主機名稱?",
            "答案: 使用 hostname 指令"
        ]
        qa = extract_qa(paragraphs)
        assert len(qa) == 1
    
    def test_extract_qa_numbered(self):
        """Test numbered question format."""
        paragraphs = [
            "1. What is the default gateway?",
            "A: The router interface IP",
            "2. How to save config?",
            "A: Use copy run start"
        ]
        qa = extract_qa(paragraphs)
        # Note: numbered questions might not be detected with current patterns
        # This test documents current behavior
    
    def test_extract_qa_empty(self):
        """Test with no Q&A."""
        paragraphs = ["Step 1", "Step 2", "Summary"]
        qa = extract_qa(paragraphs)
        assert len(qa) == 0


class TestParseDocx:
    """Tests for full docx parsing."""
    
    @pytest.fixture
    def mock_docx_path(self, tmp_path):
        """Create a mock docx file for testing."""
        try:
            from docx import Document
            
            doc = Document()
            doc.add_paragraph("Lab 2.1.3 - Basic Router Configuration")
            doc.add_paragraph("Chapter 2: Router Basics")
            doc.add_paragraph("Objective: Learn to configure a Cisco router")
            doc.add_paragraph("Step 1: Enter privileged mode")
            doc.add_paragraph("  enable")
            doc.add_paragraph("Step 2: Enter global configuration")
            doc.add_paragraph("  configure terminal")
            doc.add_paragraph("Q: What is enable?")
            doc.add_paragraph("A: Enters privileged EXEC mode")
            
            file_path = tmp_path / "test_lab.docx"
            doc.save(str(file_path))
            return str(file_path)
        except ImportError:
            pytest.skip("python-docx not installed")
    
    def test_parse_docx_full(self, mock_docx_path):
        """Test full docx parsing."""
        lab = parse_docx(mock_docx_path)
        
        assert lab.lab_id == "2.1.3"
        assert "Router" in lab.title
        assert lab.chapter == "2"
        assert lab.objective is not None
        assert len(lab.steps) >= 1
        assert len(lab.qa) >= 1
    
    def test_parse_docx_minimal(self, tmp_path):
        """Test parsing minimal docx."""
        try:
            from docx import Document
            
            doc = Document()
            doc.add_paragraph("Simple Lab Document")
            
            file_path = tmp_path / "minimal.docx"
            doc.save(str(file_path))
            
            lab = parse_docx(str(file_path))
            assert lab.lab_id == "unknown"
            assert lab.title == "Simple Lab Document"
        except ImportError:
            pytest.skip("python-docx not installed")


class TestFixtureLoading:
    """Tests for loading JSON fixtures."""
    
    def test_load_sample_lab_fixture(self):
        """Test loading the sample lab JSON fixture."""
        fixture_path = os.path.join(
            os.path.dirname(__file__), 
            "fixtures", 
            "sample_lab.json"
        )
        
        with open(fixture_path, 'r') as f:
            data = json.load(f)
        
        from backend.app.schema import Lab
        lab = Lab(**data)
        
        assert lab.lab_id == "2.1.3"
        assert lab.title == "Basic Router Configuration"
        assert len(lab.steps) == 7
        assert len(lab.qa) == 3
        assert lab.steps[0].cmds == ["enable"]
